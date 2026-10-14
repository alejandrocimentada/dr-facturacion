#!/usr/bin/env python3
"""Check Dominican invoices (RNC/cédula, NCF/e-NCF, ITBIS) and write an error report.

Usage:
    python scripts/check_invoices.py facturas.xlsx
    python scripts/check_invoices.py facturas.csv --map ncf="No. Comprobante" --map fecha="Fecha Doc"

The input file is only read, never modified. The report is written next to it as
<input-name>_reporte.xlsx (sheets: Resumen, Problemas, Facturas) and a short Spanish
summary is printed to the terminal.

Exit codes: 0 = checked (even if problems were found), 1 = file error,
2 = columns could not be detected (use --map).
"""

from __future__ import annotations

import argparse
import datetime as dt
import math
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
import yaml
from stdnum.do import cedula as std_cedula
from stdnum.do import ncf as std_ncf
from stdnum.do import rnc as std_rnc

SKILL_DIR = Path(__file__).resolve().parent.parent
RULES_DIR = SKILL_DIR / "rules"

ERROR, ADVERTENCIA, INFO = "ERROR", "ADVERTENCIA", "INFO"
SEVERITY_ORDER = {ERROR: 0, ADVERTENCIA: 1, INFO: 2}
STALE_DAYS = 90
LINE_TOLERANCE = 0.01     # RD$ per line
INVOICE_TOLERANCE = 0.05  # RD$ per invoice
EPS = 1e-9

# Every rule the checker can report: code -> (severity, short description).
RULES = {
    "R1-RNC-EMISOR": (ERROR, "RNC/cédula del emisor vacío, mal formado o con dígito verificador inválido"),
    "R1-RNC-COMPRADOR": (ERROR, "RNC/cédula del comprador mal formado o con dígito verificador inválido"),
    "R1-CEDULA-CEROS": (INFO, "Cédula sin ceros a la izquierda (guardada como número); se completó"),
    "R2-NCF-ESTRUCTURA": (ERROR, "NCF/e-NCF vacío o con estructura incorrecta"),
    "R2-NCF-TIPO": (ERROR, "Tipo de comprobante desconocido"),
    "R3-CREDITO-SIN-COMPRADOR": (ERROR, "Crédito fiscal (B01/E31) sin RNC/cédula válido del comprador"),
    "R3-CONSUMO-CON-RNC": (ADVERTENCIA, "Factura de consumo (B02/E32) emitida a un comprador con RNC"),
    "R3-EXENTO-CON-ITBIS": (ADVERTENCIA, "Régimen especial o exportación (B14/E44, B16/E46) con ITBIS cobrado"),
    "R3-NOTA-SIN-REFERENCIA": (ADVERTENCIA, "Nota de débito/crédito (B03/B04/E33/E34) sin NCF modificado"),
    "R4-NCF-DUPLICADO": (ERROR, "El mismo NCF se usa en más de una factura del archivo"),
    "R4-LINEA-REPETIDA": (ADVERTENCIA, "Línea idéntica repetida dentro de la misma factura"),
    "R4-SECUENCIA-SALTO": (INFO, "Saltos en la secuencia del mismo emisor y tipo"),
    "R4-SECUENCIA-VENCIDA": (ERROR, "Factura emitida después del vencimiento de la secuencia"),
    "R5-TASA-INVALIDA": (ERROR, "Tasa de ITBIS no reconocida"),
    "R5-NUMERO-INVALIDO": (ERROR, "Cantidad, precio, ITBIS o total vacío o ilegible"),
    "R5-ITBIS-LINEA": (ERROR, "ITBIS de la línea no coincide con cantidad × precio × tasa"),
    "R5-TOTAL-FACTURA": (ERROR, "Total de la factura no coincide con base + ITBIS"),
    "R5-TASA-REDUCIDA": (INFO, "Línea con tasa reducida (16%)"),
    "R6-FECHA-INVALIDA": (ERROR, "Fecha vacía o ilegible"),
    "R6-FECHA-FUTURA": (ERROR, "Fecha en el futuro"),
}

# Plain-language title and "what to do" for each rule (Resumen sheet and web app).
RULE_HELP = {
    "R1-RNC-EMISOR": ("El RNC o la cédula del emisor no es válido",
                      "Copia el RNC exacto desde una factura o documento oficial."),
    "R1-RNC-COMPRADOR": ("El RNC o la cédula del comprador no es válido",
                         "Pídele el RNC o la cédula correcta al cliente."),
    "R1-CEDULA-CEROS": ("Cédula sin ceros al inicio",
                        "Nada urgente. Para evitarlo, guarda esa columna como texto."),
    "R2-NCF-ESTRUCTURA": ("El NCF tiene un formato incorrecto",
                          "Copia el NCF de nuevo desde la factura original."),
    "R2-NCF-TIPO": ("Tipo de comprobante desconocido",
                    "Revisa que el NCF esté bien escrito."),
    "R3-CREDITO-SIN-COMPRADOR": ("Crédito fiscal sin RNC del comprador",
                                 "Agrega el RNC del cliente, o usa una factura de consumo (B02)."),
    "R3-CONSUMO-CON-RNC": ("Factura de consumo a una empresa",
                           "Confirma con el cliente: tal vez necesitaba crédito fiscal (B01)."),
    "R3-EXENTO-CON-ITBIS": ("ITBIS en una factura que no debería llevarlo",
                            "Revisa con tu contador si ese ITBIS va o no."),
    "R3-NOTA-SIN-REFERENCIA": ("Nota de crédito o débito sin factura original",
                               "Escribe el NCF de la factura original en la columna ncf_modificado."),
    "R4-NCF-DUPLICADO": ("NCF repetido",
                         "Revisa cuál es el correcto. Cada factura debe tener su propio NCF."),
    "R4-LINEA-REPETIDA": ("Línea repetida",
                          "Bórrala si se copió por error."),
    "R4-SECUENCIA-SALTO": ("Salto en la numeración",
                           "Si son facturas anuladas, deben ir en el formato 608 de la DGII."),
    "R4-SECUENCIA-VENCIDA": ("Secuencia vencida",
                             "Pide una secuencia nueva a la DGII y habla con tu contador."),
    "R5-TASA-INVALIDA": ("Tasa de ITBIS no válida",
                         "Corrige la tasa: 18%, 16%, 0% o exento. Lo normal es 18%."),
    "R5-NUMERO-INVALIDO": ("Número vacío o ilegible",
                           "Llena la cantidad, el precio, el ITBIS y el total con números."),
    "R5-ITBIS-LINEA": ("El ITBIS no cuadra",
                       "Recalcula el ITBIS de la línea: cantidad × precio × tasa."),
    "R5-TOTAL-FACTURA": ("El total no cuadra",
                         "Recalcula el total: monto sin ITBIS + ITBIS."),
    "R5-TASA-REDUCIDA": ("Línea con tasa del 16%",
                         "Confirma que ese producto de verdad lleva 16%."),
    "R6-FECHA-INVALIDA": ("Fecha vacía o inválida",
                          "Escribe la fecha correcta (DD/MM/AAAA)."),
    "R6-FECHA-FUTURA": ("Fecha en el futuro",
                        "Corrige el año, el mes o el día."),
}

DISCLAIMER = ("Esto no es asesoría fiscal. Revisa estos resultados con tu contador antes de "
              "presentar cualquier declaración.")

# Canonical column -> accepted header names (already normalized: lowercase, no accents,
# non-alphanumerics -> "_", "%" -> "pct").
COLUMN_ALIASES = {
    "fecha": ["fecha", "fecha_emision", "fecha_de_emision", "fecha_comprobante", "fecha_factura",
              "date", "invoice_date", "issue_date"],
    "rnc_emisor": ["rnc_emisor", "rnc_cedula_emisor", "emisor", "rnc_vendedor", "rnc_proveedor",
                   "issuer_rnc", "issuer_id", "seller_rnc", "seller_id", "supplier_rnc"],
    "rnc_comprador": ["rnc_comprador", "rnc_cedula_comprador", "cedula_comprador", "comprador",
                      "rnc_cliente", "cliente_rnc", "rnc_cedula_cliente", "buyer_rnc", "buyer_id",
                      "customer_rnc", "customer_id"],
    "ncf": ["ncf", "e_ncf", "encf", "ncf_e_ncf", "ncf_encf", "comprobante", "numero_comprobante",
            "no_comprobante", "comprobante_fiscal", "fiscal_number", "ncf_number"],
    "descripcion": ["descripcion", "concepto", "detalle", "producto", "articulo", "description",
                    "item", "product"],
    "cantidad": ["cantidad", "cant", "qty", "quantity", "units"],
    "precio_unitario": ["precio_unitario", "precio_unit", "precio", "valor_unitario", "unit_price",
                        "price"],
    "tasa_itbis": ["tasa_itbis", "tasa", "itbis_tasa", "itbis_pct", "pct_itbis", "tasa_pct",
                   "porcentaje_itbis", "itbis_rate", "tax_rate", "vat_rate", "rate"],
    "itbis": ["itbis", "monto_itbis", "itbis_facturado", "itbis_cobrado", "total_itbis",
              "tax", "tax_amount", "vat", "vat_amount", "itbis_amount"],
    "total": ["total", "monto_total", "total_linea", "importe", "importe_total", "amount",
              "line_total", "total_amount"],
    "vencimiento_secuencia": ["vencimiento_secuencia", "fecha_vencimiento_secuencia",
                              "fecha_vencimiento", "vencimiento", "valido_hasta",
                              "sequence_expiry", "sequence_expiration", "expiry", "expiration_date"],
    "ncf_modificado": ["ncf_modificado", "ncf_afectado", "ncf_referencia", "ncf_original",
                       "modified_ncf", "affected_ncf", "original_ncf", "reference_ncf"],
}
REQUIRED_COLUMNS = ["fecha", "rnc_emisor", "rnc_comprador", "ncf", "cantidad",
                    "precio_unitario", "tasa_itbis", "itbis", "total"]
OPTIONAL_COLUMNS = ["descripcion", "vencimiento_secuencia", "ncf_modificado"]

EXEMPT_WORDS = {"exento", "exenta", "exentos", "ex", "e", "exempt", "exempted", "no_aplica"}


# --------------------------------------------------------------------------- data types

@dataclass
class Issue:
    fila: int | None          # spreadsheet row number (header = row 1)
    ncf: str
    regla: str
    mensaje: str
    filas: list[int] = field(default_factory=list)  # all rows affected (for "estado")

    @property
    def severidad(self) -> str:
        return RULES[self.regla][0]


@dataclass
class NcfInfo:
    text: str
    series: str
    tipo: str
    sequence: int
    flags: dict


@dataclass
class Line:
    fila: int
    ncf_text: str
    ncf: NcfInfo | None
    emisor: str
    comprador: str
    comprador_kind: str       # vacio / rnc / cedula / invalido
    fecha: dt.date | None
    base: float | None
    itbis: float | None
    itbis_expected: float | None
    total: float | None
    ncf_modificado: str
    signature: tuple          # all mapped values, to spot repeated lines


@dataclass
class Result:
    issues: list[Issue]
    lines: list[Line]
    invoice_count: int
    stdnum_overrides: list[str]

    def counts(self) -> dict[str, int]:
        out = {ERROR: 0, ADVERTENCIA: 0, INFO: 0}
        for issue in self.issues:
            out[issue.severidad] += 1
        return out

    def row_status(self) -> dict[int, str]:
        status: dict[int, str] = {}
        for issue in self.issues:
            if issue.severidad == INFO:
                continue
            for fila in issue.filas or [issue.fila]:
                if status.get(fila) != ERROR:
                    status[fila] = issue.severidad
        return status


class ColumnMappingError(Exception):
    def __init__(self, missing: list[str], found: list[str]):
        super().__init__(f"Columnas no encontradas: {', '.join(missing)}")
        self.missing = missing
        self.found = found


# --------------------------------------------------------------------------- rules & input

def load_rules(rules_dir: Path = RULES_DIR) -> dict:
    with open(rules_dir / "itbis_rates.yaml", encoding="utf-8") as fh:
        itbis = yaml.safe_load(fh)
    with open(rules_dir / "ncf_types.yaml", encoding="utf-8") as fh:
        ncf = yaml.safe_load(fh)
    # YAML may read "01" keys as strings already; normalize defensively.
    ncf["types"] = {s: {str(k).zfill(2): (v or {}) for k, v in t.items()}
                    for s, t in ncf["types"].items()}
    return {"itbis": itbis, "ncf": ncf}


def rates_age_days(rules: dict, today: dt.date) -> int:
    verified = dt.date.fromisoformat(str(rules["itbis"]["last_verified"]))
    return (today - verified).days


def normalize_header(name) -> str:
    text = str(name).replace("%", " pct ")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def detect_columns(headers: list, manual: dict[str, str] | None = None) -> dict[str, str]:
    """Return {canonical_name: original_header}. Raises ColumnMappingError."""
    mapping: dict[str, str] = {}
    by_norm: dict[str, str] = {}
    for h in headers:
        by_norm.setdefault(normalize_header(h), h)

    for canon, header in (manual or {}).items():
        canon = normalize_header(canon)
        if canon not in COLUMN_ALIASES:
            raise ValueError(f"Campo desconocido en --map: {canon}")
        match = header if header in headers else by_norm.get(normalize_header(header))
        if match is None:
            raise ValueError(f"La columna '{header}' indicada en --map no existe en el archivo")
        mapping[canon] = match

    used = set(mapping.values())
    for canon, aliases in COLUMN_ALIASES.items():
        if canon in mapping:
            continue
        for alias in aliases:
            header = by_norm.get(alias)
            if header is not None and header not in used:
                mapping[canon] = header
                used.add(header)
                break

    missing = [c for c in REQUIRED_COLUMNS if c not in mapping]
    if missing:
        raise ColumnMappingError(missing, [str(h) for h in headers])
    return mapping


def read_table(path: Path, sheet: str | None = None) -> pd.DataFrame:
    """Read CSV/Excel into memory without touching the file. All cells kept as-is."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xls"):
        df = pd.read_excel(path, sheet_name=sheet or 0, dtype=object)
    elif suffix in (".csv", ".txt"):
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        first_line = text.splitlines()[0] if text else ""
        sep = max([",", ";", "\t", "|"], key=first_line.count)
        from io import StringIO
        df = pd.read_csv(StringIO(text), sep=sep, dtype=str, keep_default_na=False)
    else:
        raise ValueError(f"Formato no soportado: {suffix} (usa .csv o .xlsx)")
    df = df.dropna(how="all")
    df = df[~df.apply(lambda r: all(is_blank(v) for v in r), axis=1)]
    return df


# --------------------------------------------------------------------------- value parsers

def is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if value is pd.NaT:
        return True
    return str(value).strip() == ""


def clean_id(value) -> str:
    if is_blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]
    return re.sub(r"[\s\-]", "", text)


def classify_id(value) -> tuple[str, str, str]:
    """Return (kind, normalized, detail). kind: vacio / rnc / cedula / cedula_ceros / invalido."""
    text = clean_id(value)
    if not text:
        return "vacio", "", ""
    if not text.isdigit():
        return "invalido", text, "contiene caracteres que no son números"
    if len(text) == 9:
        if std_rnc.is_valid(text):
            return "rnc", text, ""
        if std_cedula.is_valid(text.zfill(11)):
            return "cedula_ceros", text.zfill(11), ""
        return "invalido", text, "el dígito verificador del RNC no es correcto"
    if len(text) == 11:
        if std_cedula.is_valid(text):
            return "cedula", text, ""
        return "invalido", text, "el dígito verificador de la cédula no es correcto"
    if len(text) < 11 and std_cedula.is_valid(text.zfill(11)):
        return "cedula_ceros", text.zfill(11), ""
    return "invalido", text, f"tiene {len(text)} dígitos (un RNC tiene 9 y una cédula 11)"


def parse_number(value) -> float | None:
    """Parse amounts like 1180, '1,180.00', 'RD$ 1.180,00'. None if blank; ValueError if garbage."""
    if is_blank(value):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = re.sub(r"(?i)rd\$|us\$|\$|\s", "", str(value))
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        if re.fullmatch(r"-?\d{1,3}(,\d{3})+", text):
            text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
    number = float(text)  # raises ValueError
    if math.isnan(number) or math.isinf(number):
        raise ValueError(value)
    return -number if negative else number


DATE_FORMATS = ["%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y",
                "%d/%m/%y", "%Y/%m/%d", "%Y%m%d"]


def parse_date(value) -> dt.date | None:
    """Day-first, as used in the Dominican Republic. None if blank; ValueError if unreadable."""
    if is_blank(value):
        return None
    if isinstance(value, (pd.Timestamp, dt.datetime)):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # Excel serial date
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(value))).date()
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"fecha ilegible: {value}")


def parse_rate(value, rates: dict) -> tuple[str, float | None]:
    """Return (rate_key, rate). rate is None for exempt. ValueError if not an allowed rate."""
    if is_blank(value):
        raise ValueError("vacía")
    word = normalize_header(value)
    if word in EXEMPT_WORDS:
        if "exento" in rates:
            return "exento", None
        raise ValueError(value)
    number = parse_number(str(value).replace("%", ""))
    if number > 1:
        number /= 100
    for key, rate in rates.items():
        if rate is not None and abs(rate - number) < 1e-6:
            return key, rate
    raise ValueError(value)


def parse_ncf(value, ncf_rules: dict) -> tuple[NcfInfo | None, str | None, str]:
    """Return (info, rule_code_if_invalid, message)."""
    text = re.sub(r"[\s\-]", "", "" if is_blank(value) else str(value)).upper()
    if not text:
        return None, "R2-NCF-ESTRUCTURA", "La línea no tiene NCF; agrega el número de comprobante."
    series_rules = ncf_rules["series"]
    series = text[0]
    if series not in series_rules:
        if len(text) == 19 and series in "AP":
            msg = (f"El NCF {text} usa el formato anterior a mayo 2018, que ya no es válido; "
                   "solicita una secuencia B o E.")
        else:
            msg = (f"El NCF {text} debe empezar con B (NCF) o E (e-NCF); "
                   "verifica que se copió correctamente.")
        return None, "R2-NCF-ESTRUCTURA", msg
    expected = series_rules[series]["length"]
    if len(text) != expected or not text[1:].isdigit():
        seq = series_rules[series]["sequence_digits"]
        return None, "R2-NCF-ESTRUCTURA", (
            f"El NCF {text} tiene {len(text)} caracteres; un {'e-NCF' if series == 'E' else 'NCF'} "
            f"debe ser {series} + 2 dígitos de tipo + {seq} dígitos de secuencia ({expected} en total).")
    tipo = text[1:3]
    types = ncf_rules["types"][series]
    if tipo not in types:
        valid = ", ".join(series + t for t in types)
        return None, "R2-NCF-TIPO", (
            f"El tipo {series}{tipo} del NCF {text} no existe; los tipos válidos son {valid}.")
    return NcfInfo(text, series, tipo, int(text[3:]), types[tipo]), None, ""


def rates_note(rules: dict) -> str:
    """Spanish note about the rates (falls back to the English `note`)."""
    itbis = rules["itbis"]
    return " ".join(str(itbis.get("nota_es") or itbis.get("note", "")).split())


def money(value: float) -> str:
    sign = "-" if value < -0.004 else ""
    return f"{sign}RD${abs(value):,.2f}"


# --------------------------------------------------------------------------- checks

def check_dataframe(df: pd.DataFrame, colmap: dict[str, str], rules: dict,
                    today: dt.date) -> Result:
    rates = rules["itbis"]["rates"]
    ncf_rules = rules["ncf"]
    issues: list[Issue] = []
    lines: list[Line] = []
    stdnum_overrides: set[str] = set()

    def get(row, canon):
        header = colmap.get(canon)
        return None if header is None else row[header]

    def add(fila, ncf, regla, mensaje, filas=None):
        issues.append(Issue(fila, ncf, regla, mensaje, filas or [fila]))

    for position, (_, row) in enumerate(df.iterrows()):
        fila = position + 2  # header is row 1
        raw_ncf = get(row, "ncf")
        ncf_text = re.sub(r"[\s\-]", "", "" if is_blank(raw_ncf) else str(raw_ncf)).upper()

        # Rule 2: NCF structure / type
        ncf, code, msg = parse_ncf(raw_ncf, ncf_rules)
        if code:
            add(fila, ncf_text, code, msg)
        elif not std_ncf.is_valid(ncf.text):
            stdnum_overrides.add(ncf.series + ncf.tipo)

        # Rule 1: RNC / cédula
        kind_e, emisor, detail = classify_id(get(row, "rnc_emisor"))
        if kind_e == "vacio":
            add(fila, ncf_text, "R1-RNC-EMISOR",
                "Falta el RNC/cédula del emisor; complétalo con el RNC de quien emite la factura.")
        elif kind_e == "invalido":
            add(fila, ncf_text, "R1-RNC-EMISOR",
                f"El RNC/cédula del emisor {emisor} no es válido: {detail}; corrígelo con el número correcto.")
        kind_c, comprador, detail = classify_id(get(row, "rnc_comprador"))
        if kind_c == "invalido":
            add(fila, ncf_text, "R1-RNC-COMPRADOR",
                f"El RNC/cédula del comprador {comprador} no es válido: {detail}; pide el número correcto al cliente.")
        for kind, number, who in ((kind_e, emisor, "emisor"), (kind_c, comprador, "comprador")):
            if kind == "cedula_ceros":
                add(fila, ncf_text, "R1-CEDULA-CEROS",
                    f"La cédula del {who} perdió los ceros iniciales; se leyó como {number}. "
                    "Guarda esa columna como texto para evitarlo.")
        kind_e = "cedula" if kind_e == "cedula_ceros" else kind_e
        kind_c = "cedula" if kind_c == "cedula_ceros" else kind_c

        # Rule 6: dates (+ rule 4: expired sequence)
        fecha = None
        try:
            fecha = parse_date(get(row, "fecha"))
            if fecha is None:
                add(fila, ncf_text, "R6-FECHA-INVALIDA", "La línea no tiene fecha; agrega la fecha de emisión.")
            elif fecha > today:
                add(fila, ncf_text, "R6-FECHA-FUTURA",
                    f"La fecha {fecha:%d/%m/%Y} es posterior a hoy; corrige la fecha de emisión.")
        except ValueError:
            add(fila, ncf_text, "R6-FECHA-INVALIDA",
                f"La fecha '{get(row, 'fecha')}' no es una fecha válida; corrígela usando el formato DD/MM/AAAA.")
        if "vencimiento_secuencia" in colmap:
            raw_venc = get(row, "vencimiento_secuencia")
            try:
                venc = parse_date(raw_venc)
            except ValueError:
                venc = None
                add(fila, ncf_text, "R6-FECHA-INVALIDA",
                    f"El vencimiento de secuencia '{raw_venc}' no es una fecha válida; corrígelo usando el formato DD/MM/AAAA.")
            if venc and fecha and fecha > venc:
                add(fila, ncf_text, "R4-SECUENCIA-VENCIDA",
                    f"La factura es del {fecha:%d/%m/%Y} pero la secuencia venció el {venc:%d/%m/%Y}; "
                    "el comprobante no es válido y debe emitirse con una secuencia vigente.")

        # Rule 5: ITBIS math
        numbers = {}
        for canon, label in (("cantidad", "la cantidad"), ("precio_unitario", "el precio unitario"),
                             ("itbis", "el ITBIS"), ("total", "el total")):
            raw = get(row, canon)
            try:
                value = parse_number(raw)
                if value is None and canon == "itbis":
                    value = 0.0  # blank ITBIS = none charged
                if value is None:
                    add(fila, ncf_text, "R5-NUMERO-INVALIDO", f"Falta {label}; complétalo para poder verificar la línea.")
                numbers[canon] = value
            except ValueError:
                add(fila, ncf_text, "R5-NUMERO-INVALIDO",
                    f"No se pudo leer {label} '{raw}'; usa solo números (ej. 1180.00).")
                numbers[canon] = None

        rate_key, rate = None, None
        raw_rate = get(row, "tasa_itbis")
        try:
            rate_key, rate = parse_rate(raw_rate, rates)
        except ValueError:
            allowed = ", ".join("exento" if r is None else f"{r * 100:g}%" for r in rates.values())
            add(fila, ncf_text, "R5-TASA-INVALIDA",
                f"La tasa de ITBIS '{'' if is_blank(raw_rate) else raw_rate}' no es válida; usa una de: {allowed}.")

        base = None
        if numbers["cantidad"] is not None and numbers["precio_unitario"] is not None:
            base = numbers["cantidad"] * numbers["precio_unitario"]
        itbis_expected = None
        if base is not None and rate_key is not None:
            itbis_expected = round(base * (rate or 0.0), 2)
            charged = numbers["itbis"]
            if charged is not None and abs(itbis_expected - charged) > LINE_TOLERANCE + EPS:
                if rate is None:
                    msg = (f"La línea es exenta pero cobra {money(charged)} de ITBIS; "
                           "elimina el ITBIS o corrige la tasa.")
                else:
                    msg = (f"El ITBIS cobrado es {money(charged)} pero debería ser {money(itbis_expected)} "
                           f"({money(base)} × {rate * 100:g}%); corrige el monto o la tasa.")
                add(fila, ncf_text, "R5-ITBIS-LINEA", msg)
        if rate_key == "reducida":
            add(fila, ncf_text, "R5-TASA-REDUCIDA",
                "Esta línea usa la tasa reducida de 16%; confirma que el producto está en la lista de "
                "tasa reducida vigente de la DGII.")

        raw_mod = get(row, "ncf_modificado")
        signature = tuple("" if is_blank(row[h]) else str(row[h]).strip() for h in colmap.values())
        lines.append(Line(fila, ncf_text, ncf, emisor, comprador, kind_c, fecha, base,
                          numbers["itbis"], itbis_expected, numbers["total"],
                          "" if is_blank(raw_mod) else str(raw_mod).strip(), signature))

    invoice_count = _check_invoices(lines, issues)
    _check_sequences(lines, issues)
    return Result(issues, lines, invoice_count, sorted(stdnum_overrides))


def _check_invoices(lines: list[Line], issues: list[Issue]) -> int:
    """Group lines into invoices and run invoice-level rules. Returns the invoice count."""
    # Lines with the same emitter + NCF belong to one invoice, unless their date or buyer
    # differ: then the NCF was used twice.
    by_ncf: dict[tuple, dict[tuple, list[Line]]] = {}
    for line in lines:
        key = (line.emisor, line.ncf_text) if line.ncf_text else ("", f"#fila{line.fila}")
        header = (line.fecha, line.comprador)
        by_ncf.setdefault(key, {}).setdefault(header, []).append(line)

    count = 0
    for (_, ncf_text), invoices in by_ncf.items():
        groups = list(invoices.values())
        first_row = groups[0][0].fila
        for group in groups[1:]:
            rows = [l.fila for l in group]
            issues.append(Issue(rows[0], ncf_text, "R4-NCF-DUPLICADO",
                                f"El NCF {ncf_text} ya se usó en la fila {first_row} para otra factura "
                                "(otra fecha o comprador); cada comprobante solo puede usarse una vez.",
                                rows))
        for group in groups:
            count += 1
            _check_one_invoice(group, issues)
    return count


def _check_one_invoice(group: list[Line], issues: list[Issue]) -> None:
    first = group[0]
    rows = [l.fila for l in group]
    ncf = first.ncf

    def add(regla, mensaje, fila=first.fila):
        issues.append(Issue(fila, first.ncf_text, regla, mensaje, rows))

    seen: dict[tuple, int] = {}
    for line in group:
        if line.signature in seen:
            issues.append(Issue(line.fila, line.ncf_text, "R4-LINEA-REPETIDA",
                                f"Esta línea es idéntica a la fila {seen[line.signature]}; "
                                "verifica que la factura no se haya copiado dos veces.", [line.fila]))
        else:
            seen[line.signature] = line.fila

    if ncf is not None:
        name = f"{ncf.series}{ncf.tipo}"
        if ncf.flags.get("requires_buyer_id") and first.comprador_kind not in ("rnc", "cedula"):
            detail = "no tiene" if first.comprador_kind == "vacio" else "tiene un"
            suffix = "" if first.comprador_kind == "vacio" else " inválido"
            add("R3-CREDITO-SIN-COMPRADOR",
                f"La factura de crédito fiscal {name} {detail} RNC/cédula del comprador{suffix}; "
                "agrega un RNC/cédula válido o emite una factura de consumo.")
        if ncf.flags.get("consumer") and first.comprador_kind == "rnc":
            credit = "E31" if ncf.series == "E" else "B01"
            add("R3-CONSUMO-CON-RNC",
                f"Factura de consumo {name} emitida a un comprador con RNC {first.comprador}: el comprador "
                f"no podrá usarla como crédito de ITBIS ni como gasto para ISR; probablemente necesitaba "
                f"una {credit}.")
        if ncf.flags.get("usually_exempt"):
            charged = [l for l in group if (l.itbis or 0) > LINE_TOLERANCE]
            if charged:
                add("R3-EXENTO-CON-ITBIS",
                    f"El comprobante {name} normalmente va exento o al 0%, pero cobra ITBIS; "
                    "confirma que corresponde cobrarlo.", charged[0].fila)
        if ncf.flags.get("modifies") and not any(l.ncf_modificado for l in group):
            add("R3-NOTA-SIN-REFERENCIA",
                f"La nota {name} no indica el NCF que modifica; agrega el NCF de la factura original "
                "en la columna ncf_modificado.")

    if all(l.base is not None and l.itbis is not None and l.total is not None for l in group):
        expected = sum(l.base for l in group) + sum(l.itbis for l in group)
        charged = sum(l.total for l in group)
        totals = {round(l.total, 2) for l in group}
        repeated_invoice_total = (len(group) > 1 and len(totals) == 1
                                  and abs(group[0].total - expected) <= INVOICE_TOLERANCE + EPS)
        if abs(charged - expected) > INVOICE_TOLERANCE + EPS and not repeated_invoice_total:
            add("R5-TOTAL-FACTURA",
                f"El total de la factura es {money(charged)} pero base + ITBIS da {money(expected)} "
                f"(diferencia {money(charged - expected)}); corrige el total o los montos de las líneas.")


def _check_sequences(lines: list[Line], issues: list[Issue]) -> None:
    sequences: dict[tuple, dict[int, Line]] = {}
    for line in lines:
        if line.ncf is not None:
            key = (line.emisor, line.ncf.series + line.ncf.tipo)
            sequences.setdefault(key, {}).setdefault(line.ncf.sequence, line)
    for (_, prefix), by_seq in sequences.items():
        numbers = sorted(by_seq)
        for prev, nxt in zip(numbers, numbers[1:]):
            missing = nxt - prev - 1
            if missing <= 0:
                continue
            line = by_seq[nxt]
            width = len(line.ncf.text) - 3
            start = f"{prefix}{prev + 1:0{width}d}"
            end = f"{prefix}{nxt - 1:0{width}d}"
            span = start if missing == 1 else f"{start} a {end}"
            issues.append(Issue(line.fila, line.ncf.text, "R4-SECUENCIA-SALTO",
                                (f"Falta 1 comprobante en la secuencia ({span}); si fue anulado, "
                                 if missing == 1 else
                                 f"Faltan {missing} comprobantes en la secuencia ({span}); si fueron anulados, ")
                                + "confirma que estén reportados en el formato 608 de la DGII.", [line.fila]))


# --------------------------------------------------------------------------- report

def sorted_issues(issues: list[Issue]) -> list[Issue]:
    return sorted(issues, key=lambda i: (SEVERITY_ORDER[i.severidad], i.fila or 0, i.regla))


def report_path_for(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_reporte.xlsx")


def itbis_totals(result: Result) -> tuple[float, float]:
    comparable = [l for l in result.lines if l.itbis is not None and l.itbis_expected is not None]
    return sum(l.itbis for l in comparable), sum(l.itbis_expected for l in comparable)


def top_actions(result: Result, n: int = 3) -> list[dict]:
    """The n most important rules to fix: most severe first, then the most affected rows."""
    by_rule: dict[str, set[int]] = {}
    for issue in result.issues:
        by_rule.setdefault(issue.regla, set()).update(f for f in (issue.filas or [issue.fila]) if f)
    ranked = sorted(by_rule, key=lambda r: (SEVERITY_ORDER[RULES[r][0]], -len(by_rule[r]), r))
    return [{"regla": r, "severidad": RULES[r][0], "titulo": RULE_HELP[r][0],
             "accion": RULE_HELP[r][1], "filas": sorted(by_rule[r])} for r in ranked[:n]]


def rows_text(filas: list[int], limit: int = 8) -> str:
    shown = ", ".join(str(f) for f in filas[:limit])
    extra = f" y {len(filas) - limit} más" if len(filas) > limit else ""
    return f"{'Fila' if len(filas) == 1 else 'Filas'} {shown}{extra}"


def _cell_value(value, is_amount: bool):
    """Original cell value for the Facturas sheet; plain numeric text in amount columns becomes a number."""
    if is_blank(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if is_amount and isinstance(value, str) and re.fullmatch(r"\s*-?\d+(\.\d+)?\s*", value):
        return float(value)
    return value


def write_report(result: Result, df: pd.DataFrame, input_path: Path, output: Path,
                 rules: dict, today: dt.date, colmap: dict[str, str] | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    # status -> (strong, light); same palette on every sheet
    palette = {"OK": ("2F6B27", "EEF8EC"), ADVERTENCIA: ("9A5B00", "FFF3DC"),
               ERROR: ("B42318", "FDECEA"), INFO: ("5B6570", "F1F3F2")}
    ink, muted, rule_color, accent = "1B1F23", "5B6570", "D9DED9", "9BE28F"
    sans, mono = "Calibri", "Consolas"
    thin = Side(style="thin", color=rule_color)

    def font(size=11, color=ink, bold=False, name=sans):
        return Font(name=name, size=size, color=color, bold=bold)

    def fill(color):
        return PatternFill("solid", fgColor=color)

    def chip(cell, estado):
        strong, light = palette[estado]
        cell.font = font(10, strong, True)
        cell.fill = fill(light)
        cell.alignment = Alignment(horizontal="center", vertical="top")

    def style_table(ws, widths, wrap_cols=(), mono_cols=(), freeze="A2"):
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = font(10, ink, True, mono)
            cell.fill = fill("F3F5F4")
            cell.border = Border(bottom=Side(style="thin", color=ink))
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 22
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_cols)
                if cell.column in mono_cols:
                    cell.font = font(10, ink, name=mono)
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()

    # 1. Resumen: header, number cards, top actions, details, notes
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    counts = result.counts()
    charged, expected = itbis_totals(result)
    diff = round(charged - expected, 2)
    age = rates_age_days(rules, today)
    status = result.row_status()
    lines_ok = sum(1 for l in result.lines if l.fila not in status)

    first, last = 2, 15                       # content spans B:O, cards are 2 columns + a gutter
    gutters = {1: 3, 4: 2, 7: 2, 10: 2, 13: 2, 16: 3}
    for col in range(1, 17):
        ws.column_dimensions[get_column_letter(col)].width = gutters.get(col, 12)
    chars_per_line = 125

    def text_height(text, size=10.5):
        lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in str(text).split("\n"))
        return lines * (size * 1.45) + 8

    def put(row, col, value, f, end=None, align="left", valign="top", wrap=False):
        if end and end > col:
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=end)
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = f
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
        return cell

    def label(row, text, col=first, end=last, color=muted):
        put(row, col, text, font(9, color, True, mono), end, valign="bottom")
        ws.row_dimensions[row].height = 18

    def spacer(row, height=14):
        ws.row_dimensions[row].height = height
        return row + 1

    label(2, "DR FACTURACIÓN  ·  REPORTE DE REVISIÓN", color=palette["OK"][0])
    put(3, first, "Revisión de facturas", font(24, ink, True), last, valign="center")
    ws.row_dimensions[3].height = 38
    for col, end, text, value in ((2, 6, "ARCHIVO REVISADO", input_path.name),
                                  (8, 9, "FECHA DE REVISIÓN", today.strftime("%d/%m/%Y")),
                                  (11, 15, "TASAS ITBIS VERIFICADAS EL",
                                   f"{rules['itbis']['last_verified']} (hace {age} días)")):
        label(4, text, col, end)
        put(5, col, value, font(11, ink), end)
    ws.row_dimensions[5].height = 20
    for col in range(first, last + 1):
        ws.cell(row=6, column=col).fill = fill(accent)
    ws.row_dimensions[6].height = 4
    r = spacer(7, 18)

    if age > STALE_DAYS:
        text = (f"ATENCIÓN: las tasas de ITBIS no se verifican desde hace {age} días; "
                "confírmalas en dgii.gov.do.")
        cell = put(r, first, text, font(10.5, palette[ERROR][0], True), last, valign="center", wrap=True)
        for col in range(first, last + 1):
            ws.cell(row=r, column=col).fill = fill(palette[ERROR][1])
        ws.row_dimensions[r].height = text_height(text) + 6
        r = spacer(r + 1)

    cards = [
        ("FACTURAS REVISADAS", result.invoice_count, ink,
         f"{len(result.lines)} líneas · {lines_ok} OK", ink),
        ("ERRORES", counts[ERROR], palette[ERROR][0] if counts[ERROR] else ink,
         "corregir antes de declarar", palette[ERROR][0]),
        ("ADVERTENCIAS", counts[ADVERTENCIA], palette[ADVERTENCIA][0] if counts[ADVERTENCIA] else ink,
         "revisar con atención", palette[ADVERTENCIA][0]),
        ("INFO", counts[INFO], ink, "solo confirmar", palette[INFO][0]),
        ("DIFERENCIA ITBIS (RD$)", diff, palette[ERROR][0] if abs(diff) > 0.004 else palette["OK"][0],
         "cobrado − recalculado", "5DBB4E"),
    ]
    for index, (title, value, color, note, top) in enumerate(cards):
        a, b = 2 + index * 3, 3 + index * 3
        label(r, title, a, b)
        number = put(r + 1, a, value, font(26, color, True), b)
        number.number_format = "#,##0.00" if isinstance(value, float) else "#,##0"
        put(r + 2, a, note, font(9.5, muted), b)
        for row in range(r, r + 3):
            for col in (a, b):
                cell = ws.cell(row=row, column=col)
                cell.border = Border(
                    left=thin if col == a else None, right=thin if col == b else None,
                    top=Side(style="medium", color=top) if row == r else None,
                    bottom=thin if row == r + 2 else None)
                cell.alignment = Alignment(horizontal="left", indent=1,
                                           vertical=("bottom", "center", "top")[row - r])
    ws.row_dimensions[r].height = 26
    ws.row_dimensions[r + 1].height = 40
    ws.row_dimensions[r + 2].height = 24
    r = spacer(r + 3, 24)

    label(r, "QUÉ HACER AHORA")
    r += 1
    actions = top_actions(result)
    for number, action in enumerate(actions, start=1):
        put(r, first, number, font(11, palette["OK"][0], True, mono), align="right").number_format = "00"
        rows_note = f"{rows_text(action['filas'])} · {action['severidad']}"
        rich = CellRichText(
            TextBlock(InlineFont(rFont=sans, sz=11, b=True, color=ink), action["titulo"] + ". "),
            TextBlock(InlineFont(rFont=sans, sz=11, color=ink), action["accion"] + "  "),
            TextBlock(InlineFont(rFont=sans, sz=10, color=palette[action["severidad"]][0]), rows_note),
        )
        put(r, first + 1, rich, font(11), last, wrap=True)
        plain = f"{action['titulo']}. {action['accion']}  {rows_note}"
        ws.row_dimensions[r].height = max(24, text_height(plain, 11))
        for col in range(first, last + 1):
            ws.cell(row=r, column=col).border = Border(bottom=thin)
        r += 1
    if not actions:
        put(r, first, "No se encontraron problemas en este archivo.", font(11, palette["OK"][0], True), last)
        ws.row_dimensions[r].height = 24
        r += 1
    r = spacer(r, 24)

    label(r, "DETALLE")
    r += 1
    for title, value in (("Líneas revisadas", len(result.lines)), ("Líneas OK", lines_ok),
                         ("ITBIS cobrado (RD$)", round(charged, 2)),
                         ("ITBIS recalculado (RD$)", round(expected, 2)), ("Diferencia (RD$)", diff)):
        put(r, first, title, font(10.5, ink), first + 3, valign="center")
        cell = put(r, first + 4, value, font(10.5, ink, True), first + 5, align="right", valign="center")
        cell.number_format = "#,##0.00" if isinstance(value, float) else "#,##0"
        for col in range(first, first + 6):
            ws.cell(row=r, column=col).border = Border(bottom=thin)
        ws.row_dimensions[r].height = 20
        r += 1
    r = spacer(r, 24)

    note = rates_note(rules)
    label(r, "NOTA SOBRE TASAS · LEY 30-26")
    put(r + 1, first, note, font(10, muted), last, wrap=True)
    ws.row_dimensions[r + 1].height = text_height(note, 10)
    r = spacer(r + 2, 10)
    label(r, "AVISO")
    put(r + 1, first, DISCLAIMER, font(10.5, ink, True), last, wrap=True)
    ws.row_dimensions[r + 1].height = text_height(DISCLAIMER)
    r += 2
    ws.print_area = f"A1:P{r}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # 2. Problemas
    ws = wb.create_sheet("Problemas")
    ws.append(["fila", "ncf", "severidad", "regla", "mensaje"])
    for issue in sorted_issues(result.issues):
        ws.append([issue.fila, issue.ncf, issue.severidad, issue.regla, issue.mensaje])
    style_table(ws, [7, 17, 15, 28, 100], wrap_cols={5}, mono_cols={2, 4})
    for row in ws.iter_rows(min_row=2):
        chip(row[2], row[2].value)

    # 3. Facturas: original data + estado
    ws = wb.create_sheet("Facturas")
    columns = [str(c) for c in df.columns]
    amount_headers = {(colmap or {}).get(c) for c in ("cantidad", "precio_unitario", "itbis", "total")}
    ws.append(["fila"] + columns + ["estado"])
    estados = []
    for position, (_, row) in enumerate(df.iterrows()):
        fila = position + 2
        estados.append(status.get(fila, "OK"))
        values = [_cell_value(v, header in amount_headers) for header, v in row.items()]
        ws.append([fila] + values + [estados[-1]])
    widths = [7]
    for col_cells in list(ws.columns)[1:]:
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells[1:])
        widths.append(min(max(longest + 3, len(str(col_cells[0].value)) + 5, 9), 38))
    style_table(ws, widths, wrap_cols={i for i, w in enumerate(widths, start=1) if w == 38},
                mono_cols={1}, freeze="B2")
    for row, estado in zip(ws.iter_rows(min_row=2), estados):
        for cell in row[:-1]:
            if estado in (ERROR, ADVERTENCIA):
                cell.fill = fill(palette[estado][1])
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
            elif isinstance(cell.value, dt.datetime):
                cell.number_format = "DD/MM/YYYY"
        chip(row[-1], estado)

    # Never overwrite the input; if the report is open in Excel, pick another name.
    candidates = [output] + [output.with_name(f"{output.stem}_{n}{output.suffix}") for n in range(1, 50)]
    for candidate in candidates:
        if candidate.resolve() == input_path.resolve():
            continue
        try:
            wb.save(candidate)
            return candidate
        except PermissionError:
            continue
    raise PermissionError(f"No se pudo guardar el reporte en {output}")


def format_summary(result: Result, input_path: Path, report: Path | None, rules: dict,
                   today: dt.date, top: int = 10) -> str:
    counts = result.counts()
    charged, expected = itbis_totals(result)
    age = rates_age_days(rules, today)
    out = [
        f"Revisión de facturas: {input_path.name}",
        f"  Facturas revisadas: {result.invoice_count} ({len(result.lines)} líneas)",
        f"  ERROR: {counts[ERROR]} | ADVERTENCIA: {counts[ADVERTENCIA]} | INFO: {counts[INFO]}",
        f"  ITBIS cobrado: {money(charged)} | recalculado: {money(expected)} | "
        f"diferencia: {money(charged - expected)}",
    ]
    issues = sorted_issues(result.issues)
    if issues:
        out.append("")
        out.append("Principales problemas:")
        for issue in issues[:top]:
            where = f"fila {issue.fila}" + (f", {issue.ncf}" if issue.ncf else "")
            out.append(f"  [{issue.severidad}] {where} ({issue.regla}): {issue.mensaje}")
        if len(issues) > top:
            out.append(f"  ... y {len(issues) - top} más en el reporte.")
    else:
        out.append("")
        out.append("No se encontraron problemas.")
    out.append("")
    if report:
        out.append(f"Reporte: {report}")
    out.append(f"Tasas de ITBIS verificadas el {rules['itbis']['last_verified']} (hace {age} días).")
    if age > STALE_DAYS:
        out.append(f"AVISO: las tasas de ITBIS no se verifican desde hace {age} días (más de "
                   f"{STALE_DAYS}); confírmalas en dgii.gov.do antes de confiar en los resultados.")
    if rates_note(rules):
        out.append("NOTA: " + rates_note(rules))
    if result.stdnum_overrides:
        out.append("NOTA: python-stdnum no reconoce los tipos " + ", ".join(result.stdnum_overrides)
                   + "; se aceptaron según rules/ncf_types.yaml.")
    return "\n".join(out)


# --------------------------------------------------------------------------- CLI

def parse_map(values: list[str]) -> dict[str, str]:
    mapping = {}
    for item in values or []:
        if "=" not in item:
            raise ValueError(f"--map debe tener la forma campo=Encabezado (recibido: {item})")
        canon, header = item.split("=", 1)
        mapping[canon.strip()] = header.strip()
    return mapping


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass

    parser = argparse.ArgumentParser(description="Revisa facturas dominicanas (RNC, NCF, ITBIS).")
    parser.add_argument("input", type=Path, help="Archivo .csv o .xlsx con una fila por línea de factura")
    parser.add_argument("--map", action="append", default=[], metavar="campo=Encabezado",
                        help="Asigna manualmente una columna (repetible), ej. --map ncf=\"No. Comprobante\"")
    parser.add_argument("--sheet", help="Hoja de Excel a leer (por defecto la primera)")
    parser.add_argument("--output", type=Path, help="Ruta del reporte (por defecto <archivo>_reporte.xlsx)")
    parser.add_argument("--today", type=dt.date.fromisoformat, default=dt.date.today(),
                        help=argparse.SUPPRESS)  # for reproducible tests
    parser.add_argument("--rules-dir", type=Path, default=RULES_DIR, help=argparse.SUPPRESS)
    args = parser.parse_args(argv)

    input_path: Path = args.input
    if not input_path.is_file():
        print(f"No se encontró el archivo: {input_path}", file=sys.stderr)
        return 1
    try:
        rules = load_rules(args.rules_dir)
        df = read_table(input_path, args.sheet)
        colmap = detect_columns(list(df.columns), parse_map(args.map))
    except ColumnMappingError as exc:
        print("No se pudieron identificar todas las columnas.", file=sys.stderr)
        print(f"  Faltan: {', '.join(exc.missing)}", file=sys.stderr)
        print(f"  Encabezados encontrados: {', '.join(exc.found)}", file=sys.stderr)
        print('  Indica a qué columna corresponde cada campo con --map, ej.: '
              '--map ncf="No. Comprobante" --map fecha="Fecha Doc"', file=sys.stderr)
        return 2
    except (ValueError, OSError) as exc:
        print(f"Error leyendo el archivo: {exc}", file=sys.stderr)
        return 1

    result = check_dataframe(df, colmap, rules, args.today)
    output = args.output or report_path_for(input_path)
    report = write_report(result, df, input_path, output, rules, args.today, colmap)
    print(format_summary(result, input_path, report, rules, args.today))
    return 0


if __name__ == "__main__":
    sys.exit(main())
