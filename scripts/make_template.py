#!/usr/bin/env python3
"""Regenerate templates/plantilla_facturas.xlsx (input template for check_invoices.py)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "templates" / "plantilla_facturas.xlsx"

# (header, required, number_format, description, example)
COLUMNS = [
    ("fecha", True, "DD/MM/YYYY", "Fecha de emisión del comprobante (DD/MM/AAAA).", "15/09/2026"),
    ("rnc_emisor", True, "@", "RNC (9 dígitos) o cédula (11 dígitos) de quien emite la factura.", "101850043"),
    ("rnc_comprador", True, "@", "RNC o cédula del comprador. Puede quedar vacío en facturas de consumo.", "130123454"),
    ("ncf", True, "@", "NCF (B + 2 + 8 = 11 caracteres) o e-NCF (E + 2 + 10 = 13 caracteres).", "B0100000001"),
    ("descripcion", False, "@", "Descripción del producto o servicio.", "Servicio de consultoría"),
    ("cantidad", True, "#,##0.00", "Cantidad vendida.", 1),
    ("precio_unitario", True, "#,##0.00", "Precio unitario SIN ITBIS (RD$).", 10000),
    ("tasa_itbis", True, "@", "18, 16, 0 o exento.", "18"),
    ("itbis", True, "#,##0.00", "ITBIS cobrado en la línea (RD$).", 1800),
    ("total", True, "#,##0.00", "Total de la línea = cantidad × precio + ITBIS (RD$).", 11800),
    ("vencimiento_secuencia", False, "DD/MM/YYYY", "Opcional: fecha de vencimiento de la secuencia de NCF.", "31/12/2026"),
    ("ncf_modificado", False, "@", "Opcional: NCF de la factura original (solo notas de débito/crédito B03/B04/E33/E34).", ""),
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    required_fill = PatternFill("solid", fgColor="263238")
    optional_fill = PatternFill("solid", fgColor="607D8B")
    for col, (name, required, fmt, desc, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = required_fill if required else optional_fill
        cell.comment = Comment(desc, "dr-facturacion")
        letter = cell.column_letter
        ws.column_dimensions[letter].width = max(len(name) + 4, 14)
        for row in range(2, 1001):
            ws.cell(row=row, column=col).number_format = fmt
    ws.freeze_panes = "A2"
    rates = DataValidation(type="list", formula1='"18,16,0,exento"', allow_blank=True)
    ws.add_data_validation(rates)
    rates.add("H2:H1000")

    info = wb.create_sheet("Instrucciones")
    info.append(["columna", "obligatoria", "descripción", "ejemplo"])
    for name, required, _, desc, example in COLUMNS:
        info.append([name, "sí" if required else "no", desc, example])
    info.append([])
    info.append(["Una fila por cada línea de factura. Las líneas de una misma factura repiten fecha, emisor, comprador y NCF."])
    info.append(["Guarda las columnas de RNC/cédula y NCF como texto para no perder los ceros iniciales."])
    info.append(["Los encabezados pueden estar en español o inglés; si no se reconocen, el asistente te pedirá asignarlos."])
    for cell in info[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = required_fill
    for width, letter in zip((24, 12, 80, 26), "ABCD"):
        info.column_dimensions[letter].width = width
    for row in info.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
